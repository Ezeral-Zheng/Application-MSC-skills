#!/usr/bin/env python3
"""Prepare a Tencent Docs spreadsheet write plan from Major List rows.

This script does the deterministic local work:
- read Postgraduate Program List/Major List.xlsx
- produce candidate matches for requested "school + program" items
- map matched program fields to a target sheet schema
- produce a safe write plan for Tencent Docs SHEET MCP tools

It does not call the Tencent Docs MCP server directly. In Codex, run this script
to create a plan, inspect matched/ambiguous/unmatched items, and let the
tencent-program-importer skill decide whether to apply the generated payload
with the MCP sheet_set_range_value tool.
"""

from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


VAULT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MAJOR_LIST = VAULT_ROOT / "Postgraduate Program List" / "Major List.xlsx"
DEFAULT_TEMPLATE = VAULT_ROOT / "Application_System" / "templates" / "学校专业信息收集模板.xlsx"
DEFAULT_OUTPUT = VAULT_ROOT / "Application_System" / "reports" / "tencent_program_import_plan.json"

MAJOR_SHEET_NAME = "Taught Masters"
TEMPLATE_SHEET_NAME = "Program_Database"

SOURCE_ALIASES = {
    "学校英文名": ["学校英文名", "School", "University", "University Name", "学校"],
    "项目中文译名": ["项目中文译名", "项目中文名", "Program Chinese Name"],
    "项目英文名": ["项目英文名", "Program Name", "专业英文名", "项目名称"],
    "学制": ["学制", "Duration"],
    "语言要求（IELTS简化/原文保留）": ["语言要求（IELTS简化/原文保留）", "语言要求", "IELTS"],
    "学费（非本地学生，数字）": ["学费（非本地学生，数字）", "学费", "Tuition"],
    "项目官网URL": ["项目官网URL", "官网", "来源URL", "URL"],
    "CompassEdu匹配专业": ["CompassEdu匹配专业", "指南者匹配专业"],
    "CompassEdu匹配URL": ["CompassEdu匹配URL", "指南者URL"],
    "CompassEdu匹配分": ["CompassEdu匹配分", "匹配分"],
    "课程（中文 / English；分号分隔）": ["课程（中文 / English；分号分隔）", "核心课程", "课程"],
    "课程抓取备注": ["课程抓取备注", "备注"],
}

TARGET_TO_SOURCE = {
    "学校英文名": "学校英文名",
    "项目中文名": "项目中文译名",
    "Program Name": "项目英文名",
    "学制": "学制",
    "语言要求": "语言要求（IELTS简化/原文保留）",
    "学费": "学费（非本地学生，数字）",
    "来源URL": "项目官网URL",
    "项目官网URL": "项目官网URL",
    "核心课程": "课程（中文 / English；分号分隔）",
    "原文摘录/证据": "课程（中文 / English；分号分隔）",
}

DERIVED_TARGETS = {
    "项目定位/一句话卖点": "derived_positioning",
    "可写入PS的匹配句": "derived_ps_sentence",
}


@dataclass
class ProgramRow:
    row_number: int
    values: dict[str, str]
    norm_key: str


def normalize(value: str) -> str:
    value = str(value or "").strip().lower()
    value = re.sub(r"[’'`]", "", value)
    value = re.sub(r"&", "and", value)
    value = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", value)
    return value


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_headers_from_template(path: Path = DEFAULT_TEMPLATE, sheet_name: str = TEMPLATE_SHEET_NAME) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
    if len(rows) < 2:
        raise SystemExit(f"Template sheet missing header row 2: {path}::{sheet_name}")
    return [clean_cell(cell) for cell in rows[1] if clean_cell(cell)]


def load_headers_arg(value: str | None) -> list[str] | None:
    if not value:
        return None
    path = Path(value)
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
    else:
        data = json.loads(value)
    if isinstance(data, dict) and "headers" in data:
        data = data["headers"]
    if not isinstance(data, list):
        raise SystemExit("--target-headers must be a JSON list or {'headers': [...]}")
    return [clean_cell(item) for item in data]


def load_major_rows(path: Path = DEFAULT_MAJOR_LIST, sheet_name: str = MAJOR_SHEET_NAME) -> list[ProgramRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    raw_headers = [clean_cell(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = canonicalize_headers(raw_headers)
    rows: list[ProgramRow] = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {headers[i]: clean_cell(row[i]) if i < len(row) else "" for i in range(len(headers))}
        school = values.get("学校英文名", "")
        program_en = values.get("项目英文名", "")
        program_zh = values.get("项目中文译名", "")
        key = " ".join(part for part in [school, program_en, program_zh] if part)
        if not normalize(key):
            continue
        rows.append(ProgramRow(row_number=row_number, values=values, norm_key=normalize(key)))
    return rows


def canonicalize_headers(headers: list[str]) -> list[str]:
    canonical: list[str] = []
    for header in headers:
        matched = None
        for target, aliases in SOURCE_ALIASES.items():
            if header in aliases:
                matched = target
                break
        canonical.append(matched or header)
    return canonical


def read_items(args: argparse.Namespace) -> list[str]:
    items: list[str] = []
    if args.items:
        items.extend(args.items)
    if args.items_file:
        path = Path(args.items_file)
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#"):
                items.append(line)
    if not items:
        raise SystemExit("Pass program requests with --item or --items-file")
    return items


def match_item(item: str, rows: list[ProgramRow], min_score: float) -> dict[str, Any]:
    query = normalize(item)
    exact = [row for row in rows if query and query in row.norm_key]
    if len(exact) == 1:
        return {"status": "matched", "item": item, "score": 1.0, "row": exact[0]}
    if len(exact) > 1:
        return {"status": "ambiguous", "item": item, "candidates": rank_candidates(query, exact, limit=8)}

    ranked = rank_candidates(query, rows, limit=8)
    if ranked and ranked[0]["score"] >= min_score:
        top_score = ranked[0]["score"]
        close = [candidate for candidate in ranked if top_score - candidate["score"] < 0.03]
        if len(close) == 1:
            return {"status": "matched", "item": item, "score": top_score, "row": close[0]["row"]}
        return {"status": "ambiguous", "item": item, "candidates": close}
    return {"status": "unmatched", "item": item, "candidates": ranked}


def rank_candidates(query: str, rows: list[ProgramRow], limit: int = 8) -> list[dict[str, Any]]:
    ranked: list[dict[str, Any]] = []
    for row in rows:
        score = difflib.SequenceMatcher(None, query, row.norm_key).ratio()
        if query and query in row.norm_key:
            score = max(score, 0.98)
        ranked.append({"score": round(score, 4), "row": row})
    ranked.sort(key=lambda item: item["score"], reverse=True)
    return ranked[:limit]


def derive_value(target_header: str, source: dict[str, str], student_name: str) -> str:
    school = source.get("学校英文名", "")
    program = source.get("项目英文名", "")
    courses = source.get("课程（中文 / English；分号分隔）", "")
    if target_header == "项目定位/一句话卖点":
        if program and courses:
            first_course = courses.split("；", 1)[0]
            return f"{program} at {school} offers coursework such as {first_course}, which can support academic and professional positioning."
        if program:
            return f"{program} at {school}"
    if target_header == "可写入PS的匹配句":
        if program and courses:
            first_course = courses.split("；", 1)[0]
            prefix = f"{student_name} can connect" if student_name else "The applicant can connect"
            return f"{prefix} prior academic or project experience with {program}'s curriculum, including {first_course}."
    return ""


def mapped_row(source: dict[str, str], target_headers: list[str], student_name: str) -> dict[str, str]:
    output: dict[str, str] = {}
    for header in target_headers:
        if not header:
            continue
        if header == "学生姓名":
            output[header] = student_name
        elif header in TARGET_TO_SOURCE:
            output[header] = source.get(TARGET_TO_SOURCE[header], "")
        elif header in DERIVED_TARGETS:
            output[header] = derive_value(header, source, student_name)
        else:
            output[header] = ""
    return output


def infer_next_row(existing_csv: str | None, default_start_row: int) -> int:
    if not existing_csv:
        return default_start_row
    rows = list(csv.reader(existing_csv.splitlines()))
    last_non_empty = -1
    for idx, row in enumerate(rows):
        if any(cell.strip() for cell in row):
            last_non_empty = idx
    return max(default_start_row, last_non_empty + 1)


def build_plan(args: argparse.Namespace) -> dict[str, Any]:
    major_rows = load_major_rows(Path(args.major_list), args.major_sheet)
    target_headers = load_headers_arg(args.target_headers) or load_headers_from_template()
    items = read_items(args)
    matches = [match_item(item, major_rows, args.min_score) for item in items]

    matched_rows = []
    unmatched = []
    ambiguous = []
    for result in matches:
        if result["status"] == "matched":
            source = result["row"].values
            matched_rows.append(
                {
                    "input": result["item"],
                    "score": result["score"],
                    "major_list_row": result["row"].row_number,
                    "school": source.get("学校英文名", ""),
                    "program": source.get("项目英文名", ""),
                    "mapped": mapped_row(source, target_headers, args.student),
                    "source": source,
                }
            )
        elif result["status"] == "ambiguous":
            ambiguous.append(serialize_candidates(result))
        else:
            unmatched.append(serialize_candidates(result))

    next_row = args.start_row
    values = []
    for offset, row in enumerate(matched_rows):
        target_row = next_row + offset
        for col, header in enumerate(target_headers):
            value = row["mapped"].get(header, "")
            if value == "":
                continue
            values.append(
                {
                    "row": target_row,
                    "col": col,
                    "value_type": "STRING",
                    "string_value": value,
                }
            )

    return {
        "file_url": args.file_url,
        "file_id": args.file_id,
        "sheet_id": args.sheet_id,
        "student": args.student,
        "target_headers": target_headers,
        "start_row": next_row,
        "matched_count": len(matched_rows),
        "unmatched_count": len(unmatched),
        "ambiguous_count": len(ambiguous),
        "matched": matched_rows,
        "unmatched": unmatched,
        "ambiguous": ambiguous,
        "mcp_sheet_set_range_value": {
            "file_id": args.file_id or "<填入腾讯文档 file_id>",
            "sheet_id": args.sheet_id or "<填入子表 sheet_id>",
            "values": values,
        },
    }


def serialize_candidates(result: dict[str, Any]) -> dict[str, Any]:
    return {
        "input": result["item"],
        "status": result["status"],
        "candidates": [
            {
                "score": candidate["score"],
                "major_list_row": candidate["row"].row_number,
                "school": candidate["row"].values.get("学校英文名", ""),
                "program": candidate["row"].values.get("项目英文名", ""),
                "program_cn": candidate["row"].values.get("项目中文译名", ""),
            }
            for candidate in result.get("candidates", [])
        ],
    }


def render_markdown(plan: dict[str, Any]) -> str:
    lines = [
        "# Tencent Program Import Plan",
        "",
        f"- Student: {plan['student'] or ''}",
        f"- Target file URL: {plan['file_url'] or ''}",
        f"- Target file ID: {plan['file_id'] or ''}",
        f"- Target sheet ID: {plan['sheet_id'] or ''}",
        f"- Matched: {plan['matched_count']}",
        f"- Ambiguous: {plan['ambiguous_count']}",
        f"- Unmatched: {plan['unmatched_count']}",
        "",
        "## Matched",
        "",
    ]
    if plan["matched"]:
        for item in plan["matched"]:
            lines.append(f"- {item['input']} -> row {item['major_list_row']}: {item['school']} / {item['program']}")
    else:
        lines.append("- None")
    lines.extend(["", "## Ambiguous", ""])
    lines.extend(render_problem_items(plan["ambiguous"]))
    lines.extend(["", "## Unmatched", ""])
    lines.extend(render_problem_items(plan["unmatched"]))
    lines.extend(["", "## MCP Payload", "", "```json", json.dumps(plan["mcp_sheet_set_range_value"], ensure_ascii=False, indent=2), "```", ""])
    return "\n".join(lines)


def render_problem_items(items: list[dict[str, Any]]) -> list[str]:
    if not items:
        return ["- None"]
    lines: list[str] = []
    for item in items:
        lines.append(f"- {item['input']}")
        for candidate in item["candidates"][:5]:
            lines.append(
                f"  - {candidate['score']}: row {candidate['major_list_row']} "
                f"{candidate['school']} / {candidate['program']}"
            )
    return lines


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file-url", help="Tencent Docs spreadsheet URL, kept in the plan for traceability")
    parser.add_argument("--file-id", help="Tencent Docs spreadsheet file_id for MCP writes")
    parser.add_argument("--sheet-id", help="Tencent Docs sheet_id for MCP writes")
    parser.add_argument("--student", default="", help="Student name to fill when target has 学生姓名")
    parser.add_argument("--item", action="append", dest="items", help="Program request, e.g. 'NUS MSc Finance'")
    parser.add_argument("--items-file", help="Text file with one program request per line")
    parser.add_argument("--major-list", default=str(DEFAULT_MAJOR_LIST), help="Path to Major List.xlsx")
    parser.add_argument("--major-sheet", default=MAJOR_SHEET_NAME, help="Sheet name in Major List.xlsx")
    parser.add_argument("--target-headers", help="JSON list/file of target headers. Defaults to template Program_Database headers")
    parser.add_argument("--start-row", type=int, default=3, help="0-based row index to begin writing; default 3 skips template/example rows")
    parser.add_argument("--min-score", type=float, default=0.72, help="Minimum fuzzy match score")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Output JSON plan path")
    parser.add_argument("--markdown", help="Optional Markdown report path")
    args = parser.parse_args()

    plan = build_plan(args)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)

    if args.markdown:
        md_path = Path(args.markdown)
        md_path.parent.mkdir(parents=True, exist_ok=True)
        md_path.write_text(render_markdown(plan), encoding="utf-8")
        print(md_path)

    if plan["ambiguous_count"] or plan["unmatched_count"]:
        print(f"Blocked: {plan['ambiguous_count']} ambiguous, {plan['unmatched_count']} unmatched. Review before writing.")
        return 2
    print(f"Plan ready: {plan['matched_count']} candidate match(es). Skill/user review is still required before writing.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
