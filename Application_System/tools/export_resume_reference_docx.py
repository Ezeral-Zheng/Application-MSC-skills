#!/usr/bin/env python3
"""Export an Application-MSC resume using the supplied compact blue reference layout."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from export_resume_docx import education_detail_rows, parse_markdown


BLUE = "1F4E79"
GREY = "666666"
BLACK = "222222"
FONT = "Arial"


def set_font(run, size=9.5, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def paragraph_format(p, before=0, after=0, line=1.0, left=0, first=0):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.left_indent = Inches(left)
    fmt.first_line_indent = Inches(first)
    fmt.widow_control = True


def add_rich(p, text, size=9.5, base_color=BLACK):
    token_re = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = p.add_run(text[pos:match.start()])
            set_font(run, size=size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = p.add_run(token[2:-2])
            set_font(run, size=size, bold=True, color=base_color)
        else:
            run = p.add_run(token[1:-1])
            set_font(run, size=size, italic=True, color=GREY)
        pos = match.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_font(run, size=size, color=base_color)


def add_rule(doc):
    p = doc.add_paragraph()
    paragraph_format(p, before=0, after=1, line=Pt(1))
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    ppr.append(borders)


def add_heading(doc, title):
    p = doc.add_paragraph()
    paragraph_format(p, before=4, after=0, line=1.0)
    run = p.add_run(title.upper())
    set_font(run, size=11.5, bold=True, color=BLUE)
    add_rule(doc)


def add_metadata(doc, title, subtitle, date=""):
    p = doc.add_paragraph()
    paragraph_format(p, before=0, after=0, line=1.0)
    run = p.add_run(title)
    set_font(run, size=9.8, bold=True)
    if subtitle:
        run = p.add_run("  |  ")
        set_font(run, size=9.8, color=GREY)
        add_rich(p, subtitle, size=9.8, base_color=GREY)
    if date:
        run = p.add_run("  |  " + date)
        set_font(run, size=9.8, italic=True, color=GREY)
    return p


def add_bullet(doc, text, size=9.5, bullet=True):
    p = doc.add_paragraph()
    paragraph_format(p, before=0, after=0, line=1.0, left=0.18 if bullet else 0, first=-0.13 if bullet else 0)
    if bullet:
        run = p.add_run("• ")
        set_font(run, size=size, color=BLACK)
    add_rich(p, text, size=size)
    return p


def normalized_title(title):
    upper = title.upper()
    if upper == "PROJECT EXPERIENCE" or upper == "OTHER":
        return "OTHER"
    if upper == "SKILLS" or upper == "COMPUTER SKILLS":
        return "COMPUTER SKILLS"
    return upper


def export(input_path: Path, output_path: Path):
    name, contact, sections = parse_markdown(input_path)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.58)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_format(p, before=0, after=1, line=1.0)
    run = p.add_run(name or input_path.stem)
    set_font(run, size=20, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_format(p, before=0, after=5, line=1.0)
    run = p.add_run(contact)
    set_font(run, size=9.5, color=GREY)

    for section in sections:
        title = normalized_title(section["title"])
        add_heading(doc, title)
        if title == "EDUCATION":
            for item in section["items"]:
                details = education_detail_rows(item.get("details", []))
                add_metadata(doc, item["title"], item.get("subtitle", ""), item.get("date", ""))
                for detail in details:
                    p = doc.add_paragraph()
                    paragraph_format(p, before=0, after=0, line=1.0)
                    add_rich(p, detail, size=9.5)
            continue
        if title == "COMPUTER SKILLS":
            for bullet in section.get("bullets", []):
                add_bullet(doc, bullet, size=9.5, bullet=False)
            for item in section["items"]:
                text = item["title"]
                if item.get("subtitle"):
                    text += ": " + item["subtitle"]
                add_bullet(doc, text, size=9.5, bullet=False)
            continue
        for item in section["items"]:
            add_metadata(doc, item["title"], item.get("subtitle", ""), item.get("date", ""))
            for bullet in item.get("bullets", []):
                add_bullet(doc, bullet)
            for detail in item.get("details", []):
                add_bullet(doc, detail)
        for bullet in section.get("bullets", []):
            add_bullet(doc, bullet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: export_resume_reference_docx.py INPUT.md OUTPUT.docx")
    export(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
